from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from doc_parser.errors import ErrorCode
from doc_parser.models import BlockRole, ParseStatus, serialize_rows
from doc_parser.normalize import RawBlock, RawParagraph, RawTable, RawUnknown


@dataclass(frozen=True, slots=True)
class DocxParseError(Exception):
    code: ErrorCode
    path: Path

    def __str__(self) -> str:
        return f"could not read DOCX file {self.path}: {self.code.value}"


def parse_docx(path: Path) -> tuple[RawBlock, ...]:
    document = _open_document(path)
    candidates = list(_iter_body_candidates(document))
    candidates.extend(_unsupported_story_candidates(document))
    return tuple(candidates)


def _open_document(path: Path) -> Document:
    try:
        return Document(path)
    except (BadZipFile, PackageNotFoundError) as error:
        raise DocxParseError(ErrorCode.INVALID_CONTAINER, path) from error
    except KeyError as error:
        raise DocxParseError(ErrorCode.MISSING_REQUIRED_PART, path) from error


def _iter_body_candidates(document: Document) -> tuple[RawBlock, ...]:
    candidates: list[RawBlock] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            candidates.append(
                RawParagraph(text=paragraph.text, role=_paragraph_role(paragraph))
            )
            if _has_unsupported_inline_content(paragraph):
                candidates.append(
                    _unsupported_candidate("paragraph contains an unsupported inline element")
                )
        elif isinstance(child, CT_Tbl):
            candidates.append(_table_candidate(Table(child, document)))
        elif child.tag.endswith("}sectPr"):
            continue
        else:
            candidates.append(_unsupported_candidate("unsupported DOCX body element"))
    return tuple(candidates)


def _paragraph_role(paragraph: Paragraph) -> BlockRole:
    style = paragraph.style
    if style is None:
        return BlockRole.BODY
    style_name = style.name
    if style_name.startswith("Heading "):
        return BlockRole.HEADING
    if style_name == "Caption":
        return BlockRole.CAPTION
    return BlockRole.BODY


def _has_unsupported_inline_content(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:object"))


def _table_candidate(table: Table) -> RawTable:
    rows = tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
    text = serialize_rows([list(row) for row in rows])
    if _is_simple_table(table, rows):
        return RawTable(text=text, rows=rows)
    return RawTable(
        text=text,
        rows=None,
        parse_status=ParseStatus.PARTIAL,
    )


def _is_simple_table(table: Table, rows: tuple[tuple[str, ...], ...]) -> bool:
    widths = {len(row) for row in rows}
    if len(widths) > 1:
        return False
    for row in table.rows:
        for cell in row.cells:
            if cell.tables:
                return False
            cell_properties = cell._tc.tcPr
            if cell_properties.gridSpan is not None or cell_properties.vMerge is not None:
                return False
    return True


def _unsupported_story_candidates(document: Document) -> tuple[RawUnknown, ...]:
    candidates: list[RawUnknown] = []
    for section in document.sections:
        for story_name, story in (
            ("header", section.header),
            ("footer", section.footer),
        ):
            if _story_has_content(story):
                candidates.append(
                    _unsupported_candidate(f"DOCX {story_name} content is outside v1 scope")
                )
    return tuple(candidates)


def _story_has_content(story: object) -> bool:
    return any(paragraph.text.strip() for paragraph in story.paragraphs) or bool(story.tables)


def _unsupported_candidate(message: str) -> RawUnknown:
    return RawUnknown(
        page=None,
        code=ErrorCode.UNSUPPORTED_ELEMENT,
        message=message,
    )
