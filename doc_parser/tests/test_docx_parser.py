from pathlib import Path

import pytest
from docx import Document as WordDocument
from docx.oxml import OxmlElement

from doc_parser.docx_parser import DocxParseError, parse_docx
from doc_parser.errors import ErrorCode
from doc_parser.models import BlockRole, ParseStatus, SourceFormat
from doc_parser.normalize import RawParagraph, RawTable, RawUnknown, normalize_document


def _save_document(document: WordDocument, path: Path) -> Path:
    document.save(path)
    return path


def test_parse_docx_preserves_interleaved_paragraph_and_table_order(tmp_path: Path) -> None:
    # Given: a document with a paragraph, table, and paragraph in that order.
    source = WordDocument()
    source.add_paragraph("First")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "123.4"
    source.add_paragraph("Last")

    # When: the DOCX is parsed and normalized.
    document = normalize_document(
        SourceFormat.DOCX, parse_docx(_save_document(source, tmp_path / "source.docx"))
    )

    # Then: retained blocks have the original cross-type order.
    assert [(block.type, block.text) for block in document.blocks] == [
        ("paragraph", "First"),
        ("table", "Metric\t123.4"),
        ("paragraph", "Last"),
    ]


def test_parse_docx_provides_simple_table_rows_and_text(tmp_path: Path) -> None:
    # Given: a simple rectangular DOCX table.
    source = WordDocument()
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Sales"
    table.cell(1, 1).text = "123.4"

    # When: the DOCX table becomes a raw candidate.
    candidate = parse_docx(_save_document(source, tmp_path / "table.docx"))[0]

    # Then: it has both positional rows and tab-separated text.
    assert isinstance(candidate, RawTable)
    assert candidate.rows == (("Item", "Value"), ("Sales", "123.4"))
    assert candidate.text == "Item\tValue\nSales\t123.4"


def test_parse_docx_keeps_empty_paragraph_for_normalization(tmp_path: Path) -> None:
    # Given: a blank paragraph before meaningful content.
    source = WordDocument()
    source.add_paragraph("")
    source.add_paragraph("Content")

    # When: candidates are parsed and normalized.
    candidates = parse_docx(_save_document(source, tmp_path / "blank.docx"))
    document = normalize_document(SourceFormat.DOCX, candidates)

    # Then: parsing retains the blank candidate while normalization omits it.
    assert isinstance(candidates[0], RawParagraph)
    assert candidates[0].text == ""
    assert [block.text for block in document.blocks] == ["Content"]


def test_parse_docx_assigns_only_unambiguous_heading_and_caption_roles(tmp_path: Path) -> None:
    # Given: built-in Heading and Caption styled paragraphs plus normal body text.
    source = WordDocument()
    source.add_paragraph("Title", style="Heading 1")
    source.add_paragraph("Figure 1", style="Caption")
    source.add_paragraph("Body")

    # When: the document is parsed.
    candidates = parse_docx(_save_document(source, tmp_path / "roles.docx"))

    # Then: only the recognized styles receive special roles.
    assert [candidate.role for candidate in candidates] == [
        BlockRole.HEADING,
        BlockRole.CAPTION,
        BlockRole.BODY,
    ]


def test_parse_docx_never_assigns_pages(tmp_path: Path) -> None:
    # Given: a document containing text and a simple table.
    source = WordDocument()
    source.add_paragraph("Text")
    source.add_table(rows=1, cols=1).cell(0, 0).text = "Value"

    # When: the document is parsed and normalized as DOCX.
    candidates = parse_docx(_save_document(source, tmp_path / "pages.docx"))
    document = normalize_document(SourceFormat.DOCX, candidates)

    # Then: no candidate or final block has an inferred page.
    assert all(candidate.page is None for candidate in candidates)
    assert all(block.page is None for block in document.blocks)


def test_parse_docx_marks_merged_table_partial(tmp_path: Path) -> None:
    # Given: a table with a horizontally merged cell.
    source = WordDocument()
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"

    # When: the table is parsed and normalized.
    document = normalize_document(
        SourceFormat.DOCX, parse_docx(_save_document(source, tmp_path / "merged.docx"))
    )
    table_block = document.blocks[0]

    # Then: no invented matrix is exposed.
    assert table_block.parse_status is ParseStatus.PARTIAL
    assert table_block.rows is None
    assert document.issues[0].code is ErrorCode.TABLE_ROWS_UNAVAILABLE


def test_parse_docx_marks_nested_table_partial(tmp_path: Path) -> None:
    # Given: a table cell that contains a nested table.
    source = WordDocument()
    outer = source.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "Outer"
    outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Inner"

    # When: the outer table is parsed and normalized.
    document = normalize_document(
        SourceFormat.DOCX, parse_docx(_save_document(source, tmp_path / "nested.docx"))
    )

    # Then: the table remains visible but its rows are unavailable.
    assert document.blocks[0].parse_status is ParseStatus.PARTIAL
    assert document.blocks[0].rows is None
    assert document.issues[0].code is ErrorCode.TABLE_ROWS_UNAVAILABLE


def test_parse_docx_records_unsupported_header_content(tmp_path: Path) -> None:
    # Given: a document with content in its header story.
    source = WordDocument()
    source.sections[0].header.paragraphs[0].text = "Confidential"

    # When: the DOCX is parsed and normalized.
    document = normalize_document(
        SourceFormat.DOCX, parse_docx(_save_document(source, tmp_path / "header.docx"))
    )

    # Then: the omitted story is visible as an unsupported element.
    assert document.blocks[0].type == "unknown"
    assert document.issues[0].code is ErrorCode.UNSUPPORTED_ELEMENT


def test_parse_docx_records_unsupported_inline_drawing(tmp_path: Path) -> None:
    # Given: a paragraph containing an unsupported drawing element.
    source = WordDocument()
    paragraph = source.add_paragraph("Figure")
    paragraph._p.append(OxmlElement("w:drawing"))

    # When: the DOCX is parsed and normalized.
    document = normalize_document(
        SourceFormat.DOCX, parse_docx(_save_document(source, tmp_path / "drawing.docx"))
    )

    # Then: the text and the unsupported element are both retained distinctly.
    assert [block.type for block in document.blocks] == ["paragraph", "unknown"]
    assert document.issues[0].code is ErrorCode.UNSUPPORTED_ELEMENT


def test_parse_docx_rejects_non_docx_container(tmp_path: Path) -> None:
    # Given: a non-ZIP file with a DOCX extension.
    path = tmp_path / "invalid.docx"
    path.write_text("not a DOCX", encoding="utf-8")

    # When: the file is parsed.
    with pytest.raises(DocxParseError) as raised:
        parse_docx(path)

    # Then: the boundary failure exposes the existing container error code.
    assert raised.value.code is ErrorCode.INVALID_CONTAINER
